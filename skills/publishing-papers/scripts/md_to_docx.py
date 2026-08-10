#!/usr/bin/env python3
"""Convert submission-side Markdown files into editable DOCX uploads.

Takes any list of .md files plus an output directory. Conservative 11pt layout,
which is what publisher submission systems expect for cover letters and
highlights. Requires python-docx.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument  # the class; `Document` is the factory
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def iter_markdown_blocks(path: Path) -> list[tuple[str, str]]:
    """Parse a small Markdown subset (headings, bullets, paragraphs) into typed blocks."""
    blocks: list[tuple[str, str]] = []
    paragraph: list[str] = []
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line:
            if paragraph:
                blocks.append(("p", " ".join(paragraph)))
                paragraph = []
            continue
        if line.startswith("#"):
            if paragraph:
                blocks.append(("p", " ".join(paragraph)))
                paragraph = []
            level = min(len(line) - len(line.lstrip("#")), 3)
            blocks.append((f"h{level}", line.lstrip("#").strip()))
        elif line.startswith("- ") or line.startswith("* "):
            if paragraph:
                blocks.append(("p", " ".join(paragraph)))
                paragraph = []
            blocks.append(("bullet", line[2:].strip()))
        else:
            paragraph.append(line)
    if paragraph:
        blocks.append(("p", " ".join(paragraph)))
    return blocks


def apply_document_style(document: DocxDocument, font: str | None = None) -> None:
    """Apply the conservative 11pt upload layout, optionally forcing one font family."""
    styles = document.styles
    for name in ("Normal", "Body Text"):
        if name in styles:
            if font:  # keep the default theme font unless a font is available everywhere
                styles[name].font.name = font
            styles[name].font.size = Pt(11)
    for name, size in (("Title", 14), ("Heading 1", 13), ("Heading 2", 12), ("Heading 3", 11)):
        if name in styles:
            if font:
                styles[name].font.name = font
            styles[name].font.size = Pt(size)


def add_block(document: DocxDocument, block_type: str, text: str) -> None:
    """Write one parsed Markdown block into the DOCX, dropping **bold** and `code` markers."""
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    if block_type == "h1":
        document.add_heading(text, level=1).alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif block_type in ("h2", "h3"):
        document.add_heading(text, level=int(block_type[1]))
    elif block_type == "bullet":
        document.add_paragraph(text, style="List Bullet")
    else:
        document.add_paragraph(text)


def markdown_to_docx(source: Path, target: Path, font: str | None = None) -> None:
    """Convert a single Markdown file into one editable DOCX."""
    document = Document()
    apply_document_style(document, font)
    for block_type, text in iter_markdown_blocks(source):
        add_block(document, block_type, text)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(target))


def main() -> None:
    """CLI entry point: convert each .md into a same-named .docx under --out."""
    parser = argparse.ArgumentParser(description="Convert submission Markdown to DOCX")
    parser.add_argument("sources", nargs="+", type=Path, help="one or more .md files")
    parser.add_argument("--out", type=Path, required=True, help="output directory")
    parser.add_argument(
        "--font",
        default=None,
        help="font family to force, e.g. Arial; omit to keep the Word default",
    )
    args = parser.parse_args()
    for src in args.sources:
        target = args.out / (src.stem + ".docx")
        markdown_to_docx(src, target, args.font)
        print(f"  {src}  ->  {target}")


if __name__ == "__main__":
    main()
